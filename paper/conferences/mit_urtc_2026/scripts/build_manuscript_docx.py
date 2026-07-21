#!/usr/bin/env python3
"""Build a clean Word-compatible MIT URTC manuscript DOCX from audited Markdown."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[4]
SOURCE = ROOT / "paper/conferences/mit_urtc_2026/draft/MANUSCRIPT.md"
OUTPUT = ROOT / "paper/conferences/mit_urtc_2026/draft/MIT_URTC_MANUSCRIPT.docx"
TEMPLATE = (
    ROOT
    / "paper/conferences/mit_urtc_2026/guidelines/official_paper_template_letter_transitional.docx"
)

FONT = "Times New Roman"
BODY_PT = 10
COLUMN_WIDTH_IN = 3.30


def dxa(inches: float) -> str:
    return str(round(inches * 1440))


def set_run_font(run, size: float = BODY_PT, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT
    r_pr = run._element.get_or_add_rPr()
    r_pr.rFonts.set(qn("w:ascii"), FONT)
    r_pr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def format_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE


def keep_with_next(paragraph) -> None:
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def keep_together(paragraph) -> None:
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepLines"))


def set_columns(section, count: int) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), "360" if count == 2 else "720")


def configure_section(section, *, columns: int) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(0.63)
    section.right_margin = Inches(0.63)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    set_columns(section, columns)


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    for style_name in (
        "Normal",
        "paper title",
        "Author",
        "Affiliation",
        "Abstract",
        "Keywords",
        "Heading 1",
        "Heading 2",
        "Heading 5",
        "Body Text",
        "figure caption",
        "references",
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        r_fonts.set(qn("w:ascii"), FONT)
        r_fonts.set(qn("w:hAnsi"), FONT)

    doc.styles["Normal"].font.size = Pt(BODY_PT)
    doc.styles["paper title"].font.size = Pt(24)
    doc.styles["Author"].font.size = Pt(11)
    doc.styles["Affiliation"].font.size = Pt(BODY_PT)
    for style_name in ("Abstract", "Keywords", "figure caption", "references"):
        doc.styles[style_name].font.size = Pt(BODY_PT)

    body = doc.styles["Body Text"]
    body.font.size = Pt(BODY_PT)
    body.paragraph_format.first_line_indent = Inches(0.2)
    body.paragraph_format.space_after = Pt(0.5)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    references = doc.styles["references"]
    references.paragraph_format.left_indent = Inches(0.18)
    references.paragraph_format.first_line_indent = Inches(-0.18)
    references.paragraph_format.space_after = Pt(1)
    references.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_inline(paragraph, text: str, *, size: float = BODY_PT) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size, italic=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size)


def set_cell_margin(cell, *, top=40, start=55, bottom=40, end=55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), dxa(inches))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), dxa(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), dxa(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margin(cell)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "808080")


def add_title_block(doc: Document, title: str, author: str, affiliation: str, contact: str) -> None:
    p = doc.add_paragraph(style="paper title")
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    set_run_font(p.add_run(title), 24)

    p = doc.add_paragraph(style="Author")
    set_run_font(p.add_run(author), 11)
    for value in (affiliation, contact):
        p = doc.add_paragraph(style="Affiliation")
        set_run_font(p.add_run(value), BODY_PT)


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    style_name = "Abstract" if label.startswith("Abstract") else "Keywords"
    p = doc.add_paragraph(style=style_name)
    format_paragraph(p, after=4)
    set_run_font(p.add_run(label), 10, bold=True, italic=True)
    add_inline(p, text)


def add_heading(doc: Document, text: str, level: int) -> None:
    if level == 2:
        style_name = "Heading 5" if text == "References" else "Heading 1"
        clean_text = re.sub(r"^[IVX]+\.\s+", "", text)
        p = doc.add_paragraph(style=style_name)
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=5, after=1)
        set_run_font(p.add_run(clean_text), 10)
    else:
        clean_text = re.sub(r"^[A-Z]\.\s+", "", text)
        p = doc.add_paragraph(style="Heading 2")
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=3)
        set_run_font(p.add_run(clean_text), 10, italic=True)
    keep_with_next(p)


def add_body_paragraph(doc: Document, lines: Iterable[str]) -> None:
    text = " ".join(line.strip() for line in lines).strip()
    if not text:
        return
    p = doc.add_paragraph(style="Body Text")
    format_paragraph(p)
    p.paragraph_format.first_line_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(0.5)
    add_inline(p, text)


def add_list(doc: Document, items: list[str], *, ordered: bool) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Body Text")
        format_paragraph(p)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        prefix = f"{index}. " if ordered else "- "
        set_run_font(p.add_run(prefix), BODY_PT)
        add_inline(p, item)
        if len(items) <= 4 and index < len(items):
            keep_with_next(p)


def add_reference_paragraph(doc: Document, lines: Iterable[str]) -> None:
    text = " ".join(line.strip() for line in lines).strip()
    if not text:
        return
    text = re.sub(r"^\[\d+\]\s*", "", text)
    p = doc.add_paragraph(style="references")
    format_paragraph(p)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(1)
    add_inline(p, text)


def add_table(doc: Document, caption: str, lines: list[str]) -> None:
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=1)
    set_run_font(p.add_run(caption.upper()), 10)
    keep_with_next(p)

    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    header, body = rows[0], rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [1.85, 0.62, 0.83] if len(header) == 3 else [1.24, 0.49, 0.49, 0.49, 0.59]
    set_table_geometry(table, widths)
    set_table_borders(table)

    for col, value in enumerate(header):
        cell = table.rows[0].cells[col]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, value.replace("R-squared", "R2"))
        for run in p.runs:
            run.bold = True

    for values in body:
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            set_cell_width(cell, widths[col])
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT if col == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            add_inline(p, value.replace("**", ""))
    for row in table.rows:
        row_pr = row._tr.get_or_add_trPr()
        row_pr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph()


def add_image(doc: Document, alt: str, relative_path: str) -> None:
    image_path = (SOURCE.parent / relative_path).resolve()
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2)
    p.add_run().add_picture(str(image_path), width=Inches(COLUMN_WIDTH_IN))
    keep_with_next(p)
    p = doc.add_paragraph(style="figure caption")
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    set_run_font(p.add_run(alt), 10, italic=True)
    keep_together(p)


def collect_section(lines: list[str], heading: str) -> list[str]:
    start = lines.index(heading) + 1
    out = []
    for line in lines[start:]:
        if line.startswith("## "):
            break
        out.append(line)
    return out


def parse_metadata(lines: list[str]) -> tuple[str, str, str, str]:
    title = lines[0].removeprefix("# ").strip()
    metadata = {}
    for line in lines[1:8]:
        match = re.match(r"\*\*(Author|Affiliation|Contact):\*\*\s*(.+?)\s*$", line)
        if match:
            metadata[match.group(1)] = match.group(2)
    return title, metadata["Author"], metadata["Affiliation"], metadata["Contact"]


def parse_abstract(lines: list[str]) -> tuple[str, str]:
    section = collect_section(lines, "## Abstract")
    abstract_lines, keyword_lines = [], []
    in_keywords = False
    for line in section:
        if line.startswith("**Keywords:**"):
            in_keywords = True
            keyword_lines.append(line.split("**Keywords:**", 1)[1].strip())
        elif in_keywords:
            keyword_lines.append(line.strip())
        else:
            abstract_lines.append(line.strip())
    return " ".join(filter(None, abstract_lines)), " ".join(filter(None, keyword_lines))


def build() -> None:
    lines = SOURCE.read_text().splitlines()
    title, author, affiliation, contact = parse_metadata(lines)
    abstract, keywords = parse_abstract(lines)

    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Missing MIT URTC template working copy: {TEMPLATE}")
    doc = Document(TEMPLATE)
    clear_template_body(doc)
    configure_styles(doc)
    configure_section(doc.sections[0], columns=1)
    add_title_block(doc, title, author, affiliation, contact)
    add_labeled_paragraph(doc, "Abstract: ", abstract)
    add_labeled_paragraph(doc, "Keywords: ", keywords)

    configure_section(doc.add_section(WD_SECTION.CONTINUOUS), columns=2)
    body_lines = lines[lines.index("## I. Introduction") :]
    pending: list[str] = []
    pending_caption = ""
    in_references = False
    i = 0

    def flush() -> None:
        nonlocal pending
        if in_references:
            add_reference_paragraph(doc, pending)
        else:
            add_body_paragraph(doc, pending)
        pending = []

    while i < len(body_lines):
        line = body_lines[i]
        if line.startswith("## "):
            flush()
            heading_text = line[3:].strip()
            in_references = heading_text == "References"
            add_heading(doc, heading_text, 2)
            i += 1
        elif line.startswith("### "):
            flush()
            add_heading(doc, line[4:].strip(), 3)
            i += 1
        elif line.startswith("**Table ") and line.endswith("**"):
            flush()
            pending_caption = line.strip("*")
            i += 1
        elif line.startswith("|"):
            flush()
            table_lines = []
            while i < len(body_lines) and body_lines[i].startswith("|"):
                table_lines.append(body_lines[i])
                i += 1
            add_table(doc, pending_caption, table_lines)
            pending_caption = ""
        elif match := re.match(r"^!\[(.+)\]\((.+)\)$", line):
            flush()
            add_image(doc, match.group(1), match.group(2))
            i += 1
        elif re.match(r"^\d+\.\s+", line):
            flush()
            items = []
            while i < len(body_lines) and re.match(r"^\d+\.\s+", body_lines[i]):
                text = re.sub(r"^\d+\.\s+", "", body_lines[i]).strip()
                i += 1
                while i < len(body_lines) and body_lines[i].startswith("   "):
                    text += " " + body_lines[i].strip()
                    i += 1
                items.append(text)
            add_list(doc, items, ordered=True)
        elif line.startswith("- "):
            flush()
            items = []
            while i < len(body_lines) and body_lines[i].startswith("- "):
                text = body_lines[i][2:].strip()
                i += 1
                while i < len(body_lines) and body_lines[i].startswith("  "):
                    text += " " + body_lines[i].strip()
                    i += 1
                items.append(text.rstrip(";"))
            add_list(doc, items, ordered=False)
        elif not line.strip():
            flush()
            i += 1
        else:
            pending.append(line)
            i += 1
    flush()

    props = doc.core_properties
    props.title = title
    props.author = author
    props.subject = "MIT URTC technical paper manuscript"
    props.keywords = keywords
    props.comments = "Generated from audited repository evidence."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
